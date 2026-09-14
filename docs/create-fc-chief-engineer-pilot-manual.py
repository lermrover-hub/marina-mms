from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "07-fc-chief-engineer-pilot-manual.md"
OUTPUT = ROOT / "07-fc-chief-engineer-pilot-manual.docx"

BLACK = "000000"
NAVY = "17324D"
TEAL = "13988F"
PALE = "EAF6F5"
LIGHT = "F4F6F8"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def set_run_font(run, name="Arial", size=10.5, bold=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def style_paragraph(paragraph, size=10.5, bold=False, color=BLACK):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12


def add_markdown_runs(paragraph, text):
    parts = text.split("**")
    for index, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part.replace("`", ""))
        set_run_font(run, bold=index % 2 == 1)


def add_title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_pr = title._p.get_or_add_pPr()
    title_border = title_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_pr.remove(title_border)
    run = title.add_run("คู่มือทดลองใช้ Marina MMS")
    set_run_font(run, size=28, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("สำหรับ FC และ Chief Engineer")
    set_run_font(run, size=20, bold=True, color=BLACK)
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Vercel Preview และ Supabase Staging\nฉบับ 1.0 วันที่ 14 กันยายน 2026")
    set_run_font(run, size=11, color="4B5563")
    doc.add_page_break()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value.strip()
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(paragraph, size=8.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value.strip().replace("`", "")
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                style_paragraph(paragraph, size=8.5)
                if len(value) < 18:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for style_name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style_ppr = style._element.get_or_add_pPr()
        style_border = style_ppr.find(qn("w:pBdr"))
        if style_border is not None:
            style_ppr.remove(style_border)

    add_title_page(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    first_heading_skipped = False
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# ") and not first_heading_skipped:
            first_heading_skipped = True
            index += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(line[3:], style="Heading 1")
            style_paragraph(paragraph, size=18, bold=True)
            index += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(line[4:], style="Heading 2")
            style_paragraph(paragraph, size=14, bold=True)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            if len(parsed) >= 2:
                add_table(doc, parsed[0], parsed[2:])
            continue
        if line.startswith("- [ ] "):
            paragraph = doc.add_paragraph()
            add_markdown_runs(paragraph, "☐ " + line[6:])
            style_paragraph(paragraph)
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.first_line_indent = Cm(-0.5)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_markdown_runs(paragraph, line[2:])
            style_paragraph(paragraph)
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            paragraph = doc.add_paragraph()
            add_markdown_runs(paragraph, line)
            style_paragraph(paragraph)
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.first_line_indent = Cm(-0.6)
            index += 1
            continue
        if line.startswith("**ระบบ:") or line.startswith("**สภาพแวดล้อม:") or line.startswith("**ฉบับ:") or line.startswith("**ผู้ใช้:"):
            index += 1
            continue
        paragraph = doc.add_paragraph()
        add_markdown_runs(paragraph, line.replace("  ", " "))
        style_paragraph(paragraph)
        index += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Marina MMS Pilot Manual  |  Internal Training Use")
    set_run_font(run, size=8, color="647076")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
